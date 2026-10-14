import json
import openpyxl
import sys

with open('database/semena_data.json', 'r', encoding='utf-8') as f:
    d = json.load(f)

sys.stdout.buffer.write(b'=== TRUCK_ROWS in semena_data.json ===\n')
sys.stdout.buffer.write(b'First 3 (Top of table):\n')
for t in d['truck_rows'][:3]:
    s = f"  N {t['num']} | {t['date']} | {t['supplier']} | {t['car_num']} | {t['delivered_t']} t\n"
    sys.stdout.buffer.write(s.encode('utf-8'))

sys.stdout.buffer.write(b'Last 3 (Bottom of table):\n')
for t in d['truck_rows'][-3:]:
    s = f"  N {t['num']} | {t['date']} | {t['supplier']} | {t['car_num']} | {t['delivered_t']} t\n"
    sys.stdout.buffer.write(s.encode('utf-8'))

sys.stdout.buffer.write(b'\n=== HISTORY_ROWS in semena_data.json ===\n')
sys.stdout.buffer.write(b'First 3 (Top of table):\n')
for h in d['history_rows'][:3]:
    s = f"  {h['dt']} | {h['supplier']} | {h['spec']} | {h['category']}\n"
    sys.stdout.buffer.write(s.encode('utf-8'))

wb = openpyxl.load_workbook('database/Лист_погодження.xlsx', data_only=True)
sys.stdout.buffer.write(b'\n=== Excel Sheet: Reestr_priyomok ===\n')
ws1 = wb['Реєстр_прийомок']
for r in range(1, 5):
    vals = [str(ws1.cell(r, c).value or '') for c in range(1, 7)]
    s = f"Row {r}: " + ' | '.join(vals) + '\n'
    sys.stdout.buffer.write(s.encode('utf-8'))

sys.stdout.buffer.write(b'\n=== Excel Sheet: Istoriya_zmin ===\n')
ws2 = wb['Історія_змін']
for r in range(1, 5):
    vals = [str(ws2.cell(r, c).value or '') for c in range(1, 6)]
    s = f"Row {r}: " + ' | '.join(vals) + '\n'
    sys.stdout.buffer.write(s.encode('utf-8'))
