import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re
import json
import os
import datetime

def normalize_date(d_str):
    s = str(d_str or '').strip().split()[0] if str(d_str or '').strip() else ''
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{2,4})$', s)
    if m:
        y = ('20' + m.group(3)) if len(m.group(3)) == 2 else m.group(3)
        return f"{m.group(1)}.{m.group(2)}.{y}"
    return s

def date_sort_key(d_str):
    s = normalize_date(d_str)
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{4})$', s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return '9999-99-99'

def ukr_weekday(date_str):
    s = normalize_date(date_str)
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{4})$', s)
    if m:
        dt = datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        days = {0: 'Пн', 1: 'Вт', 2: 'Ср', 3: 'Чт', 4: 'Пт', 5: 'Сб', 6: 'Нд'}
        return days.get(dt.weekday(), '')
    return ''

def extract_code_and_spec_num(text):
    s = str(text or '').strip()
    if not s or re.match(r'^\d+(\.\d+)?$', s):
        return {'code': '', 'sp_num': '', 'display': ''}
    
    sp_num = ''
    m = re.search(r'(?:Сп\.?|СП\.?|№)\s*(\d+)', s, re.I)
    if m:
        sp_num = 'Сп.' + m.group(1)
        
    code = ''
    cm = re.search(r'([A-Za-zА-Яа-яІіЇїЄєҐґ]\s*[-–]\s*\d+)', s)
    if cm:
        code = re.sub(r'\s+', '', cm.group(1).replace('–', '-')).upper()
    else:
        pm = re.search(r'\(([^)]+)\)', s)
        if pm and not re.search(r'(?:Сп\.?|СП\.?|№)\s*\d+', pm.group(1), re.I):
            code = pm.group(1).strip().upper()
        else:
            parts = s.split('(')
            candidate = parts[0].strip().upper()
            if not re.match(r'^\d+$', candidate):
                code = candidate
    if not code or re.match(r'^\d+$', code):
        return {'code': '', 'sp_num': '', 'display': ''}
    
    display = f'{code} ({sp_num})' if (sp_num and code) else code
    return {'code': code, 'sp_num': sp_num, 'display': display}

def parse_receipts(prihod_path):
    wb = openpyxl.load_workbook(prihod_path, data_only=True)
    ws = wb.active
    
    # 1. Header datetime
    dt_info = {'iso': datetime.datetime.now().strftime('%Y-%m-%dT%H:%M:%S'), 'display': datetime.datetime.now().strftime('%d.%m.%Y час %H:%M')}
    for r in range(1, min(10, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            v = str(ws.cell(r, c).value or '').strip()
            m = re.search(r'[-–—]\s*(\d{2})\.(\d{2})\.(\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?', v)
            if m:
                d, mo, y, h, mi = m.group(1), m.group(2), m.group(3), m.group(4).zfill(2), m.group(5).zfill(2)
                sec = m.group(6).zfill(2) if m.group(6) else '00'
                dt_info = {
                    'iso': f'{y}-{mo}-{d}T{h}:{mi}:{sec}',
                    'display': f'{d}.{mo}.{y} час {h}:{mi}'
                }
                break
        if dt_info['iso'] != datetime.datetime.now().strftime('%Y-%m-%dT%H:%M:%S'):
            break

    vsego_row = None
    for r in range(1, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '').strip()
        c2 = str(ws.cell(r, 2).value or '').strip()
        if c1 == 'Всего' or c2 == 'Всего':
            vsego_row = r
            break

    max_truck_row = (vsego_row - 1) if vsego_row else ws.max_row
    truck_rows = []
    for r in range(7, max_truck_row + 1):
        v1 = ws.cell(r, 1).value
        if v1 is None or not str(v1).strip().isdigit():
            continue
        num_pp = int(str(v1).strip())
        auto = str(ws.cell(r, 2).value or '').strip()
        date_raw = str(ws.cell(r, 3).value or '').strip()
        date_in = normalize_date(date_raw)
        doc_no = str(ws.cell(r, 5).value or '').strip()
        supplier = str(ws.cell(r, 8).value or '').strip()
        terms = str(ws.cell(r, 10).value or '').strip()
        spec_raw = str(ws.cell(r, 12).value or '').strip()
        
        try: netto = float(ws.cell(r, 24).value or 0)
        except: netto = 0.0
        try: total_kg = float(ws.cell(r, 25).value or 0)
        except: total_kg = 0.0
        try: zachet = float(ws.cell(r, 26).value or 0)
        except: zachet = 0.0
        
        try: w_fact = float(ws.cell(r, 32).value) if ws.cell(r, 32).value is not None else None
        except: w_fact = None
        try: s_fact = float(ws.cell(r, 35).value) if ws.cell(r, 35).value is not None else None
        except: s_fact = None
        try: oil_fact = float(ws.cell(r, 38).value) if ws.cell(r, 38).value is not None else None
        except: oil_fact = None
        
        ext = extract_code_and_spec_num(spec_raw)
        deliv_kg = netto if netto > 0 else total_kg
        
        truck_rows.append({
            'num_pp': num_pp,
            'auto': auto,
            'date': date_in,
            'doc_no': doc_no,
            'supplier': supplier,
            'terms': terms,
            'spec_raw': spec_raw,
            'spec_display': ext['display'],
            'code': ext['code'],
            'delivered_kg': deliv_kg,
            'zachet_kg': zachet,
            'diff_kg': round(deliv_kg - zachet, 3),
            'w_fact': w_fact,
            's_fact': s_fact,
            'oil_fact': oil_fact
        })

    summary_rows = []
    if vsego_row:
        for r in range(vsego_row + 1, ws.max_row + 1):
            spec_raw = str(ws.cell(r, 12).value or '').strip()
            if not spec_raw:
                continue
            ext = extract_code_and_spec_num(spec_raw)
            if not ext['code']:
                continue
            supplier = str(ws.cell(r, 9).value or ws.cell(r, 8).value or '').strip()
            terms = str(ws.cell(r, 11).value or '').strip()
            try: deliv_kg = float(ws.cell(r, 25).value or 0)
            except: deliv_kg = 0.0
            try: zach_kg = float(ws.cell(r, 26).value or 0)
            except: zach_kg = 0.0
            oil_c = ws.cell(r, 37).value if ws.cell(r, 37).value is not None else ws.cell(r, 38).value
            try: oil_fact = float(oil_c) if oil_c is not None else None
            except: oil_fact = None
            
            summary_rows.append({
                'supplier': supplier,
                'terms': terms,
                'spec_raw': spec_raw,
                'spec_display': ext['display'],
                'code': ext['code'],
                'delivered_kg': deliv_kg,
                'zachet_kg': zach_kg,
                'diff_kg': round(deliv_kg - zach_kg, 3),
                'oil_fact': oil_fact
            })

    return {
        'receipts_datetime': dt_info['display'],
        'last_receipts_iso': dt_info['iso'],
        'truck_rows': truck_rows,
        'summary_rows': summary_rows
    }

def parse_specifications(spec_path):
    wb = openpyxl.load_workbook(spec_path, data_only=True)
    ws = wb.active
    
    specs = {}
    for r in range(9, ws.max_row + 1):
        spec_cell = str(ws.cell(r, 13).value or '').strip()
        if not spec_cell or spec_cell.lower().startswith('итого'):
            continue
        ext = extract_code_and_spec_num(spec_cell)
        code = ext['code']
        if not code:
            continue
        
        date_in = normalize_date(ws.cell(r, 1).value)
        supplier = str(ws.cell(r, 4).value or '').strip()
        terms = str(ws.cell(r, 5).value or '').strip()
        manager = str(ws.cell(r, 14).value or '').strip()
        
        try: dog_tons = float(ws.cell(r, 7).value or 0)
        except: dog_tons = 0.0
        try: tons = float(ws.cell(r, 29).value or 0)
        except: tons = 0.0
        contract_tons = tons if tons > 0 else dog_tons
        
        try: dost_cost = float(ws.cell(r, 11).value or 0)
        except: dost_cost = 0.0
        try: p_post = float(ws.cell(r, 23).value or 0)
        except: p_post = 0.0
        try: p_dost = float(ws.cell(r, 26).value or 0)
        except: p_dost = 0.0
        try: p_dog = float(ws.cell(r, 27).value or 0)
        except: p_dog = 0.0
        
        price_dost = p_post if p_post > 0 else (p_dost if p_dost > 0 else (p_dog + dost_cost if p_dog > 0 else None))
        price_no_dost = p_dog if p_dog > 0 else (price_dost if (terms.upper() == 'DAP' or dost_cost == 0) else (price_dost - dost_cost if price_dost is not None else None))
        
        try: base_oil = float(ws.cell(r, 35).value or 48.0)
        except: base_oil = 48.0
        try: fact_oil = float(ws.cell(r, 34).value) if ws.cell(r, 34).value is not None else None
        except: fact_oil = None
        try: fact_w = float(ws.cell(r, 33).value) if ws.cell(r, 33).value is not None else None
        except: fact_w = None
        try: fact_s = float(ws.cell(r, 46).value) if ws.cell(r, 46).value is not None else None
        except: fact_s = None
        
        specs[code] = {
            'code': code,
            'sp_num': ext['sp_num'],
            'spec_display': ext['display'],
            'spec_full_name': spec_cell,
            'supplier': supplier,
            'delivery_terms': terms,
            'delivery_cost': dost_cost,
            'manager': manager,
            'date': date_in,
            'contract_tons': contract_tons,
            'price_delivery': price_dost,
            'price_no_delivery': price_no_dost,
            'base_oil': base_oil,
            'fact_oil': fact_oil,
            'fact_w': fact_w,
            'fact_s': fact_s
        }
    return specs

def process_datasets(prihod_data, specs_db, history_path, semena_data_path):
    history_data = {'events': [], 'spec_price_notes': {}}
    if os.path.exists(history_path):
        with open(history_path, 'r', encoding='utf-8') as f:
            try: history_data = json.load(f)
            except: pass
            
    existing_semena = {}
    if os.path.exists(semena_data_path):
        with open(semena_data_path, 'r', encoding='utf-8') as f:
            try: existing_semena = json.load(f)
            except: pass
            
    existing_by_code = {}
    for er in existing_semena.get('agg_rows', []):
        c = er.get('code') or extract_code_and_spec_num(er.get('spec', '')).get('code', '')
        if c:
            existing_by_code[c] = er
            
    truck_rows = prihod_data.get('truck_rows', [])
    summary_rows = prihod_data.get('summary_rows', [])
    
    trucks_by_code = {}
    for t in truck_rows:
        c = t.get('code')
        if c:
            trucks_by_code.setdefault(c, []).append(t)
            
    summary_by_code = {}
    for s in summary_rows:
        c = s.get('code')
        if c:
            summary_by_code[c] = s
            
    all_codes = []
    for s in summary_rows:
        c = s.get('code')
        if c and c not in all_codes: all_codes.append(c)
    for t in truck_rows:
        c = t.get('code')
        if c and c not in all_codes: all_codes.append(c)
    for c in specs_db.keys():
        if c not in all_codes: all_codes.append(c)
        
    now_ts = prihod_data.get('receipts_datetime') or datetime.datetime.now().strftime('%d.%m.%Y %H:%M')
    short_ts = now_ts.split()[0] if now_ts else datetime.datetime.now().strftime('%d.%m')
    
    agg_rows = []
    for code in all_codes:
        spec_info = specs_db.get(code, {})
        sum_info = summary_by_code.get(code, {})
        t_list = trucks_by_code.get(code, [])
        
        supplier = sum_info.get('supplier') or (t_list[0].get('supplier') if t_list else '') or spec_info.get('supplier', '')
        spec_display = spec_info.get('spec_display') or sum_info.get('spec_display') or (t_list[0].get('spec_display') if t_list else code)
        
        date_str = normalize_date(spec_info.get('date', ''))
        if not date_str and t_list:
            dates = [normalize_date(x.get('date', '')) for x in t_list if x.get('date')]
            if dates:
                dates.sort(key=date_sort_key)
                date_str = dates[0]
                
        manager = spec_info.get('manager', '')
        terms = sum_info.get('terms') or (t_list[0].get('terms') if t_list else '') or spec_info.get('delivery_terms', '')
        
        if sum_info.get('delivered_kg', 0) > 0:
            weight_delivered = sum_info['delivered_kg'] / 1000.0
            weight_zachet = sum_info.get('zachet_kg', 0) / 1000.0
        elif t_list:
            weight_delivered = sum(x.get('delivered_kg', 0) for x in t_list) / 1000.0
            weight_zachet = sum(x.get('zachet_kg', 0) for x in t_list) / 1000.0
        elif spec_info.get('contract_tons'):
            weight_delivered = 0.0
            weight_zachet = 0.0
        else:
            weight_delivered = 0.0
            weight_zachet = 0.0
            
        weight_diff = round(weight_delivered - weight_zachet, 3)
        price_delivery = spec_info.get('price_delivery')
        price_no_delivery = spec_info.get('price_no_delivery')
        base_oil = spec_info.get('base_oil', 48.0) or 48.0
        
        fact_w = None
        fact_s = None
        fact_oil = sum_info.get('oil_fact')
        
        if t_list:
            tot_w = 0.0
            w_sum, w_cnt = 0.0, 0
            s_sum, s_cnt = 0.0, 0
            o_sum, o_cnt = 0.0, 0
            for t in t_list:
                wgt = t.get('zachet_kg', 0) or t.get('delivered_kg', 0) or 1.0
                tot_w += wgt
                if t.get('w_fact') is not None:
                    w_sum += t['w_fact'] * wgt; w_cnt += 1
                if t.get('s_fact') is not None:
                    s_sum += t['s_fact'] * wgt; s_cnt += 1
                if t.get('oil_fact') is not None:
                    o_sum += t['oil_fact'] * wgt; o_cnt += 1
            if w_cnt > 0 and tot_w > 0: fact_w = w_sum / tot_w
            if s_cnt > 0 and tot_w > 0: fact_s = s_sum / tot_w
            if fact_oil is None and o_cnt > 0 and tot_w > 0: fact_oil = o_sum / tot_w
            
        if fact_w is None and spec_info.get('fact_w') is not None: fact_w = float(spec_info['fact_w'])
        if fact_s is None and spec_info.get('fact_s') is not None: fact_s = float(spec_info['fact_s'])
        if fact_oil is None and spec_info.get('fact_oil') is not None: fact_oil = float(spec_info['fact_oil'])
        
        fact_w = round(fact_w, 2) if fact_w is not None else 0.0
        fact_s = round(fact_s, 2) if fact_s is not None else 0.0
        fact_oil = round(fact_oil, 2) if fact_oil is not None else 0.0
        oil_diff = round(fact_oil - base_oil, 2) if (fact_oil > 0 and base_oil > 0) else 0.0
        
        # Change audit
        note = 'Без змін'
        change_type = 'same'
        saved_price_info = history_data.get('spec_price_notes', {}).get(code)
        
        if existing_by_code:
            if code not in existing_by_code:
                note = 'Новий запис'
                change_type = 'new'
            else:
                old = existing_by_code[code]
                price_changes = []
                other_changes = []
                old_p_deliv = old.get('price_delivered') or old.get('price_with_deliv')
                old_p_exw = old.get('price_exw') or old.get('price_no_deliv')
                old_deliv = float(old.get('delivered', 0))
                old_zach = float(old.get('credit_weight', 0))
                old_oil = float(old.get('oil_fact', 0))
                
                if price_delivery and old_p_deliv and abs(price_delivery - old_p_deliv) > 0.05:
                    dp = price_delivery - old_p_deliv
                    sign = '+' if dp >= 0 else ''
                    price_changes.append(f"Ціна з дост: {old_p_deliv:,.2f} -> {price_delivery:,.2f} ({sign}{dp:.2f})")
                if price_no_delivery and old_p_exw and abs(price_no_delivery - old_p_exw) > 0.05:
                    dp = price_no_delivery - old_p_exw
                    sign = '+' if dp >= 0 else ''
                    price_changes.append(f"Ціна без дост: {old_p_exw:,.2f} -> {price_no_delivery:,.2f} ({sign}{dp:.2f})")
                if abs(weight_delivered - old_deliv) > 0.05:
                    other_changes.append(f"Завезено: {old_deliv:.2f} -> {weight_delivered:.2f} т")
                if abs(weight_zachet - old_zach) > 0.05:
                    other_changes.append(f"Залікова: {old_zach:.2f} -> {weight_zachet:.2f} т")
                if fact_oil > 0 and old_oil > 0 and abs(fact_oil - old_oil) > 0.01:
                    other_changes.append(f"Факт масла: {old_oil:.2f}% -> {fact_oil:.2f}%")
                    
                all_cur = price_changes + other_changes
                if all_cur:
                    note = 'Змінено: ' + '; '.join(all_cur)
                    change_type = 'changed'
                elif saved_price_info and saved_price_info.get('price_note'):
                    note = f"💰 Зміна ціни: {saved_price_info['price_note']}"
                else:
                    note = 'Без змін'
                    
        agg_rows.append({
            'supplier': supplier,
            'spec': spec_display,
            'specification': spec_display,
            'code': code,
            'date': date_str,
            'spec_date': date_str,
            'manager': manager,
            'terms': terms,
            'delivery_terms': terms,
            'contract': round(float(spec_info.get('contract_tons', 0)), 2),
            'contract_t': round(float(spec_info.get('contract_tons', 0)), 2),
            'contract_tons': round(float(spec_info.get('contract_tons', 0)), 2),
            'delivered': round(weight_delivered, 3),
            'delivered_t': round(weight_delivered, 3),
            'weight_delivered': round(weight_delivered, 3),
            'credit_weight': round(weight_zachet, 3),
            'accounting_t': round(weight_zachet, 3),
            'weight_zachet': round(weight_zachet, 3),
            'weight_diff': weight_diff,
            'price_delivered': round(float(price_delivery or 0), 2),
            'price_with_deliv': round(float(price_delivery or 0), 2),
            'price_delivery': round(float(price_delivery or 0), 2),
            'price_exw': round(float(price_no_delivery or 0), 2),
            'price_no_deliv': round(float(price_no_delivery or 0), 2),
            'price_no_delivery': round(float(price_no_delivery or 0), 2),
            'oil_base': round(base_oil, 2),
            'moisture': fact_w,
            'w_fact': fact_w,
            'trash': fact_s,
            's_fact': fact_s,
            'oil_fact': fact_oil,
            'oil_diff': oil_diff,
            'note': '' if note == 'Без змін' else note,
            'change_type': change_type
        })
        
    agg_rows.sort(key=lambda x: (date_sort_key(x['date']), x['supplier']))
    for idx, r in enumerate(agg_rows):
        r['num'] = idx + 1
        
    # 2. Unfulfilled specifications
    unfulfilled_rows = []
    unf_all_rows = []
    for code, s_info in specs_db.items():
        contract_tons = float(s_info.get('contract_tons', 0))
        sum_info = summary_by_code.get(code, {})
        t_list = trucks_by_code.get(code, [])
        
        if sum_info.get('delivered_kg', 0) > 0:
            deliv_tons = sum_info['delivered_kg'] / 1000.0
            zach_tons = sum_info.get('zachet_kg', 0) / 1000.0
        elif t_list:
            deliv_tons = sum(x.get('delivered_kg', 0) for x in t_list) / 1000.0
            zach_tons = sum(x.get('zachet_kg', 0) for x in t_list) / 1000.0
        else:
            deliv_tons = 0.0
            zach_tons = 0.0
            
        rem_tons = round(max(0.0, contract_tons - deliv_tons), 3)
        pct = round((deliv_tons / contract_tons) * 100, 1) if contract_tons > 0 else 0.0
        truck_cnt = len(t_list)
        tol_10pct = round(0.10 * contract_tons, 3)
        is_fulfilled = (deliv_tons >= contract_tons) or (rem_tons <= 25.0) or (contract_tons > 0 and rem_tons <= tol_10pct)
        
        if deliv_tons <= 0.001:
            cat = 'not_started'
            status_code = 'pending'
            status_lbl = '🔴 Не приїхало (0%)'
            note_txt = 'Жодної машини не надходило'
        elif not is_fulfilled:
            cat = 'in_progress'
            status_code = 'in_progress'
            status_lbl = '🟡 У стадії виконання'
            note_txt = f"Завезено {truck_cnt} авто ({deliv_tons:.2f} т з {contract_tons:.2f} т, залишок {rem_tons:.2f} т)"
        else:
            cat = 'fulfilled'
            status_code = 'fulfilled'
            status_lbl = '🟢 Виконано'
            note_txt = f"Виконано ({truck_cnt} авто, {deliv_tons:.2f} т з {contract_tons:.2f} т)"
            
        item = {
            'category': cat,
            'status': status_lbl,
            'status_code': status_code,
            'status_raw': status_lbl,
            'status_label': status_lbl,
            'supplier': s_info.get('supplier', ''),
            'spec': s_info.get('spec_display', code),
            'spec_display': s_info.get('spec_display', code),
            'date': normalize_date(s_info.get('date', '')),
            'spec_date': normalize_date(s_info.get('date', '')),
            'manager': s_info.get('manager', ''),
            'terms': s_info.get('delivery_terms', ''),
            'delivery_terms': s_info.get('delivery_terms', ''),
            'contract': round(contract_tons, 2),
            'contract_t': round(contract_tons, 2),
            'contract_tons': round(contract_tons, 2),
            'delivered': round(deliv_tons, 3),
            'delivered_t': round(deliv_tons, 3),
            'delivered_tons': round(deliv_tons, 3),
            'credit_weight': round(zach_tons, 3),
            'zachet_tons': round(zach_tons, 3),
            'remaining': rem_tons,
            'remaining_t': rem_tons,
            'remaining_tons': rem_tons,
            'pct': pct,
            'pct_fulfilled': pct,
            'price_delivered': round(float(s_info.get('price_delivery') or 0), 2),
            'price_with_deliv': round(float(s_info.get('price_delivery') or 0), 2),
            'price_exw': round(float(s_info.get('price_no_delivery') or 0), 2),
            'note': note_txt
        }
        unf_all_rows.append(item)
        if cat != 'fulfilled':
            unfulfilled_rows.append(item)
            
    order_map = {'not_started': 0, 'in_progress': 1, 'fulfilled': 2}
    unfulfilled_rows.sort(key=lambda x: (order_map.get(x['category'], 1), x['pct'], date_sort_key(x['date'])))
    unf_all_rows.sort(key=lambda x: (order_map.get(x['category'], 1), x['pct'], date_sort_key(x['date'])))
    for idx, u in enumerate(unfulfilled_rows): u['num'] = idx + 1
    for idx, u in enumerate(unf_all_rows): u['num'] = idx + 1
    
    # 3. Daily rows
    daily_map = {}
    for t in truck_rows:
        d = normalize_date(t.get('date', '')) or 'Невідомо'
        if d not in daily_map:
            daily_map[d] = {
                'date': d,
                'trucks_count': 0,
                'delivered_kg': 0.0,
                'zachet_kg': 0.0,
                'oil_w': 0.0,
                'w_w': 0.0,
                's_w': 0.0
            }
        daily_map[d]['trucks_count'] += 1
        daily_map[d]['delivered_kg'] += float(t.get('delivered_kg', 0))
        daily_map[d]['zachet_kg'] += float(t.get('zachet_kg', 0))
        z = float(t.get('zachet_kg', 0))
        if t.get('oil_fact') is not None and z > 0: daily_map[d]['oil_w'] += t['oil_fact'] * z
        if t.get('w_fact') is not None and z > 0: daily_map[d]['w_w'] += t['w_fact'] * z
        if t.get('s_fact') is not None and z > 0: daily_map[d]['s_w'] += t['s_fact'] * z
        
    days_sorted = sorted(daily_map.keys(), key=date_sort_key)
    daily_rows = []
    for idx, d in enumerate(days_sorted):
        info = daily_map[d]
        deliv_t = round(info['delivered_kg'] / 1000.0, 3)
        zach_t = round(info['zachet_kg'] / 1000.0, 3)
        diff_t = round(deliv_t - zach_t, 3)
        diff_p = round((diff_t / deliv_t) * 100, 2) if deliv_t > 0 else 0.0
        z = info['zachet_kg']
        avg_oil = round(info['oil_w'] / z, 2) if z > 0 else 0.0
        avg_w = round(info['w_w'] / z, 2) if z > 0 else 0.0
        avg_s = round(info['s_w'] / z, 2) if z > 0 else 0.0
        dow = ukr_weekday(d)
        
        daily_rows.append({
            'num': idx + 1,
            'date': d,
            'weekday': dow,
            'day_name': dow,
            'day_of_week': dow,
            'trucks': info['trucks_count'],
            'cars': info['trucks_count'],
            'trucks_count': info['trucks_count'],
            'delivered': deliv_t,
            'delivered_t': deliv_t,
            'delivered_tons': deliv_t,
            'credit_weight': zach_t,
            'accounting_t': zach_t,
            'zachet_tons': zach_t,
            'discount': diff_t,
            'diff_t': diff_t,
            'diff_tons': diff_t,
            'discount_pct': diff_p,
            'diff_pct': diff_p,
            'oil': avg_oil,
            'oil_fact': avg_oil,
            'moisture': avg_w,
            'w_fact': avg_w,
            'trash': avg_s,
            's_fact': avg_s
        })
        
    # 4. Discount rows
    sups_map = {}
    for t in truck_rows:
        sup = t.get('supplier') or 'Без назви'
        code = t.get('code') or ''
        if sup not in sups_map:
            sups_map[sup] = {
                'supplier': sup,
                'codes': set(),
                'trucks_count': 0,
                'delivered_kg': 0.0,
                'zachet_kg': 0.0,
                'oil_w': 0.0,
                'w_w': 0.0,
                's_w': 0.0
            }
        if code: sups_map[sup]['codes'].add(code)
        sups_map[sup]['trucks_count'] += 1
        sups_map[sup]['delivered_kg'] += float(t.get('delivered_kg', 0))
        sups_map[sup]['zachet_kg'] += float(t.get('zachet_kg', 0))
        z = float(t.get('zachet_kg', 0))
        if t.get('oil_fact') is not None and z > 0: sups_map[sup]['oil_w'] += t['oil_fact'] * z
        if t.get('w_fact') is not None and z > 0: sups_map[sup]['w_w'] += t['w_fact'] * z
        if t.get('s_fact') is not None and z > 0: sups_map[sup]['s_w'] += t['s_fact'] * z
        
    sup_list = list(sups_map.values())
    sup_list.sort(key=lambda s: (s['delivered_kg'] - s['zachet_kg']), reverse=True)
    discount_rows = []
    for idx, s in enumerate(sup_list):
        deliv_t = round(s['delivered_kg'] / 1000.0, 3)
        zach_t = round(s['zachet_kg'] / 1000.0, 3)
        diff_t = round(deliv_t - zach_t, 3)
        diff_p = round((diff_t / deliv_t) * 100, 2) if deliv_t > 0 else 0.0
        z = s['zachet_kg']
        discount_rows.append({
            'num': idx + 1,
            'supplier': s['supplier'],
            'specs_count': len(s['codes']),
            'trucks': s['trucks_count'],
            'cars': s['trucks_count'],
            'trucks_count': s['trucks_count'],
            'delivered': deliv_t,
            'delivered_t': deliv_t,
            'credit_weight': zach_t,
            'accounting_t': zach_t,
            'discount': diff_t,
            'diff_t': diff_t,
            'discount_pct': diff_p,
            'diff_pct': diff_p,
            'moisture': round(s['w_w'] / z, 2) if z > 0 else 0.0,
            'trash': round(s['s_w'] / z, 2) if z > 0 else 0.0,
            'oil': round(s['oil_w'] / z, 2) if z > 0 else 0.0
        })
        
    # 5. Formatted truck rows
    formatted_trucks = []
    for idx, t in enumerate(truck_rows):
        code = t.get('code') or ''
        spec_info = specs_db.get(code, {})
        deliv_t = round(float(t.get('delivered_kg', 0)) / 1000.0, 3)
        zach_t = round(float(t.get('zachet_kg', 0)) / 1000.0, 3)
        diff_t = round(deliv_t - zach_t, 3)
        base_oil = float(spec_info.get('base_oil', 48.0) or 48.0)
        oil_fact = round(float(t['oil_fact']), 2) if t.get('oil_fact') is not None else None
        oil_diff = round(oil_fact - base_oil, 2) if (oil_fact is not None and oil_fact > 0) else None
        
        formatted_trucks.append({
            'num': idx + 1,
            'supplier': t.get('supplier') or spec_info.get('supplier', ''),
            'spec': spec_info.get('spec_display') or t.get('spec_display') or code,
            'car_num': t.get('auto', ''),
            'ttn': t.get('doc_no', ''),
            'date': normalize_date(t.get('date', '')),
            'manager': spec_info.get('manager', ''),
            'terms': t.get('terms') or spec_info.get('delivery_terms', ''),
            'delivered_t': deliv_t,
            'accounting_t': zach_t,
            'weight_diff': diff_t,
            'price_with_deliv': round(float(spec_info['price_delivery']), 2) if spec_info.get('price_delivery') else None,
            'price_no_deliv': round(float(spec_info['price_no_delivery']), 2) if spec_info.get('price_no_delivery') else None,
            'oil_base': base_oil,
            'moisture': round(float(t['w_fact']), 2) if t.get('w_fact') is not None else None,
            'trash': round(float(t['s_fact']), 2) if t.get('s_fact') is not None else None,
            'oil_fact': oil_fact,
            'oil_diff': oil_diff
        })

    # Summary calculations
    tot_specs = len(specs_db)
    unique_sups = len(set(s.get('supplier', '') for s in specs_db.values() if s.get('supplier')))
    tot_contract = round(sum(float(s.get('contract_tons', 0)) for s in specs_db.values()), 2)
    tot_deliv = round(sum(float(t.get('delivered_kg', 0)) for t in truck_rows) / 1000.0, 2)
    tot_acc = round(sum(float(t.get('zachet_kg', 0)) for t in truck_rows) / 1000.0, 2)
    tot_diff = round(tot_deliv - tot_acc, 2)
    tot_trucks = len(truck_rows)
    not_started = sum(1 for u in unf_all_rows if u['category'] == 'not_started')
    in_prog = sum(1 for u in unf_all_rows if u['category'] == 'in_progress')
    fulfilled = sum(1 for u in unf_all_rows if u['category'] == 'fulfilled')
    unf_rem_t = round(sum(u['remaining'] for u in unfulfilled_rows), 2)
    
    # Weighted quality across all trucks
    z_sum = sum(float(t.get('zachet_kg', 0)) for t in truck_rows)
    avg_oil = round(sum(t['oil_fact'] * t['zachet_kg'] for t in truck_rows if t.get('oil_fact') is not None and t.get('zachet_kg')) / z_sum, 2) if z_sum > 0 else 49.3
    avg_moist = round(sum(t['w_fact'] * t['zachet_kg'] for t in truck_rows if t.get('w_fact') is not None and t.get('zachet_kg')) / z_sum, 2) if z_sum > 0 else 7.0
    avg_trash = round(sum(t['s_fact'] * t['zachet_kg'] for t in truck_rows if t.get('s_fact') is not None and t.get('zachet_kg')) / z_sum, 2) if z_sum > 0 else 3.2
    
    summary = {
        'total_specs': tot_specs,
        'unique_suppliers': unique_sups,
        'total_contract_t': tot_contract,
        'total_delivered_t': tot_deliv,
        'total_accounting_t': tot_acc,
        'total_diff_t': tot_diff,
        'total_trucks': tot_trucks,
        'not_started_count': not_started,
        'in_progress_count': in_prog,
        'fulfilled_count': fulfilled,
        'unfulfilled_rem_t': unf_rem_t,
        'avg_oil_pct': avg_oil,
        'avg_moisture_pct': avg_moist,
        'avg_trash_pct': avg_trash,
        'deviations_count': sum(1 for r in agg_rows if r.get('note'))
    }
    
    # Update history
    history_data['last_receipts_datetime'] = prihod_data['receipts_datetime']
    history_data['last_receipts_iso'] = prihod_data['last_receipts_iso']
    
    updated_at = prihod_data['receipts_datetime']
    final_data = {
        'updated_at': updated_at,
        'report_datetime': updated_at,
        'last_receipts_iso': prihod_data['last_receipts_iso'],
        'source_subtitle': f"Дані прийомок 1С станом на {updated_at}. Зведено автоматично з 2 файлів (Приход + Специфікація).",
        'summary': summary,
        'agg_rows': agg_rows,
        'unfulfilled_rows': unfulfilled_rows,
        'unf_rows': unf_all_rows,
        'daily_rows': daily_rows,
        'discount_rows': discount_rows,
        'disc_rows': discount_rows,
        'truck_rows': formatted_trucks,
        'history_events': history_data.get('events', [])[-200:],
        'history_rows': [
            {'dt': e.get('timestamp', ''), 'supplier': e.get('supplier', ''), 'spec': e.get('specification', ''), 'manager': e.get('manager', ''), 'category': e.get('category_label', ''), 'details': e.get('summary', '')}
            for e in reversed(history_data.get('events', [])[-200:])
        ]
    }
    
    return final_data, history_data

def build_excel_workbook(data, out_path):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    hdr_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D4D4D4'),
        right=Side(style='thin', color='D4D4D4'),
        top=Side(style='thin', color='D4D4D4'),
        bottom=Side(style='thin', color='D4D4D4')
    )
    
    def add_table_sheet(title, headers, rows, col_types=None):
        ws = wb.create_sheet(title=title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(1, c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = align_center
            
        for r_data in rows:
            ws.append(r_data)
            
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(headers) + 1):
                cell = ws.cell(r, c)
                cell.border = thin_border
                ctype = col_types[c-1] if col_types and c-1 < len(col_types) else 'str'
                if ctype == 'num':
                    cell.alignment = align_right
                elif ctype == 'center':
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                    
        # Set reasonable col widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    # 1. Зведена_відомість
    h1 = ['№', 'Постачальник', 'Специфікація', 'Дата', 'Менеджер', 'Умови', 'Контракт, т', 'Завезено, т', 'Залікова, т', 'Скидка, т', 'Ціна з дост, грн', 'Ціна без дост, грн', 'База %', 'Волога %', 'Сор %', 'Олія факт %', 'Відхил %', 'Примітки']
    t1 = 'num,str,str,center,str,center,num,num,num,num,num,num,num,num,num,num,num,str'.split(',')
    r1 = [[r['num'], r['supplier'], r['spec'], r['date'], r['manager'], r['terms'], r['contract_t'], r['delivered_t'], r['accounting_t'], r['weight_diff'], r['price_with_deliv'], r['price_no_deliv'], r['oil_base'], r['moisture'], r['trash'], r['oil_fact'], r['oil_diff'], r['note']] for r in data['agg_rows']]
    add_table_sheet('Зведена_відомість', h1, r1, t1)

    # 2. Графік_завезення
    h2 = ['№', 'Дата', 'День', 'Авто', 'Фізична вага, т', 'Залікова вага, т', 'Скидка, т', 'Скидка %', 'Олійність %', 'Вологість %', 'Смітна домішка %']
    t2 = 'num,center,center,num,num,num,num,num,num,num,num'.split(',')
    r2 = [[r['num'], r['date'], r['weekday'], r['cars'], r['delivered_t'], r['accounting_t'], r['diff_t'], r['diff_pct'], r['oil'], r['moisture'], r['trash']] for r in data['daily_rows']]
    add_table_sheet('Графік_завезення', h2, r2, t2)

    # 3. Натуральна_скидка
    h3 = ['№', 'Постачальник', 'Специфікацій', 'Авто', 'Фізична вага, т', 'Залікова вага, т', 'Скидка, т', 'Скидка %', 'Волога %', 'Сор %', 'Олійність %']
    t3 = 'num,str,num,num,num,num,num,num,num,num,num'.split(',')
    r3 = [[r['num'], r['supplier'], r['specs_count'], r['cars'], r['delivered_t'], r['accounting_t'], r['discount'], r['discount_pct'], r['moisture'], r['trash'], r['oil']] for r in data['discount_rows']]
    add_table_sheet('Натуральна_скидка', h3, r3, t3)

    # 4. Невиконані_специфікації
    h4 = ['№', 'Статус', 'Постачальник', 'Специфікація', 'Дата', 'Менеджер', 'Умови', 'Контракт, т', 'Завезено, т', 'Залікова, т', 'Залишок, т', 'Виконання %', 'Ціна з дост', 'Ціна без дост', 'Коментар']
    t4 = 'num,center,str,str,center,str,center,num,num,num,num,num,num,num,str'.split(',')
    r4 = [[r['num'], r['status_label'], r['supplier'], r['spec'], r['date'], r['manager'], r['terms'], r['contract_t'], r['delivered_t'], r['credit_weight'], r['remaining_t'], r['pct'], r['price_with_deliv'], r['price_exw'], r['note']] for r in data['unfulfilled_rows']]
    add_table_sheet('Невиконані_специфікації', h4, r4, t4)

    # 5. Всі_специфікації
    r5 = [[r['num'], r['status_label'], r['supplier'], r['spec'], r['date'], r['manager'], r['terms'], r['contract_t'], r['delivered_t'], r['credit_weight'], r['remaining_t'], r['pct'], r['price_with_deliv'], r['price_exw'], r['note']] for r in data['unf_rows']]
    add_table_sheet('Всі_специфікації', h4, r5, t4)

    # 6. Реєстр_прийомок
    h6 = ['№', 'Постачальник', 'Специфікація', 'Номер авто', '№ ТТН', 'Дата прийомки', 'Менеджер', 'Умови поставки', 'Фіз. вага (нетто), т', 'Залікова вага, т', 'Скидка у вазі, т', 'Ціна з дост., грн', 'Ціна без дост., грн', 'Базова олія, %', 'Волога факт, %', 'Сор факт, %', 'Олія факт, %', 'Відхилення олії, %']
    t6 = 'num,str,str,center,center,center,str,center,num,num,num,num,num,num,num,num,num,num'.split(',')
    r6 = [[r['num'], r['supplier'], r['spec'], r['car_num'], r['ttn'], r['date'], r['manager'], r['terms'], r['delivered_t'], r['accounting_t'], r['weight_diff'], r['price_with_deliv'], r['price_no_deliv'], r['oil_base'], r['moisture'], r['trash'], r['oil_fact'], r['oil_diff']] for r in data['truck_rows']]
    add_table_sheet('Реєстр_прийомок', h6, r6, t6)

    # 7. Історія_змін
    h7 = ['Дата та час', 'Постачальник', 'Специфікація', 'Менеджер', 'Категорія', 'Опис змін']
    t7 = 'center,str,str,str,center,str'.split(',')
    r7 = [[r['dt'], r['supplier'], r['spec'], r['manager'], r['category'], r['details']] for r in data.get('history_rows', [])]
    add_table_sheet('Історія_змін', h7, r7, t7)

    wb.save(out_path)
    print(f"Excel workbook successfully written to {out_path}")

def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    prihod_path = r'C:\Users\SB1\.gemini\antigravity\brain\cf7c72ec-7c34-4571-be33-671eefad0219\.user_uploaded\media_1792006815108_027f127a.xlsx'
    spec_path = r'C:\Users\SB1\.gemini\antigravity\brain\cf7c72ec-7c34-4571-be33-671eefad0219\.user_uploaded\media_1792006814974_5eb4f28c.xlsx'
    
    db_dir = os.path.join(base_dir, 'database')
    history_path = os.path.join(db_dir, 'changes_history.json')
    semena_data_path = os.path.join(db_dir, 'semena_data.json')
    raw_prihod_path = os.path.join(db_dir, 'semena_raw_prihod.json')
    raw_spec_path = os.path.join(db_dir, 'semena_raw_spec.json')
    xlsx_out_path = os.path.join(db_dir, 'Лист_погодження.xlsx')
    
    print("1. Parsing raw 1C files...")
    prihod_data = parse_receipts(prihod_path)
    specs_db = parse_specifications(spec_path)
    
    print(f"Parsed {len(prihod_data['truck_rows'])} trucks up to {prihod_data['receipts_datetime']}")
    print(f"Parsed {len(specs_db)} specifications")
    
    print("2. Processing datasets and updating changes history...")
    final_data, history_data = process_datasets(prihod_data, specs_db, history_path, semena_data_path)
    
    print("3. Saving JSON database files...")
    with open(raw_prihod_path, 'w', encoding='utf-8') as f:
        json.dump(prihod_data, f, ensure_ascii=False, indent=2)
    with open(raw_spec_path, 'w', encoding='utf-8') as f:
        json.dump(specs_db, f, ensure_ascii=False, indent=2)
    with open(history_path, 'w', encoding='utf-8') as f:
        json.dump(history_data, f, ensure_ascii=False, indent=2)
    with open(semena_data_path, 'w', encoding='utf-8') as f:
        json.dump(final_data, f, ensure_ascii=False, indent=2)
        
    print("4. Generating Лист_погодження.xlsx...")
    build_excel_workbook(final_data, xlsx_out_path)
    
    print("\n=== SUMMARY METRICS ===")
    for k, v in final_data['summary'].items():
        print(f"  {k}: {v}")
    print(f"Total days in daily_rows: {len(final_data['daily_rows'])}")
    last_d = final_data['daily_rows'][-1]
    print(f"Last day in daily_rows: {last_d['date']} -> {last_d['delivered']} t, {last_d['cars']} cars")

if __name__ == '__main__':
    main()
